"""Tests for the evidence pipeline (task A1): extraction, indexing and structured
proposals. Gemini is never called: tests that need AI output patch the pipeline's
get_client with a fake."""

import io
import json
from types import SimpleNamespace

import docx
import pymupdf
import pytest
from google.genai import types
from sqlalchemy import select

from app.core.storage import absolute_path, save_bytes
from app.models import (
    AiRun,
    CaseObject,
    Event,
    Evidence,
    EvidenceChunk,
    EvidenceLink,
    Location,
    Observation,
    Person,
    Relationship,
)
from app.services.ai import calls, pipeline
from app.services.ai.pipeline import process_evidence
from app.services.ai.proposals import PassageIndex
from app.services.search.indexing import index_evidence
from tests.factories import make_case, make_evidence, make_object, make_person

STATEMENT = """Statement of Anita Desai, recorded 4 October 2026.

I was waiting at the Ring Road and 5th Avenue junction at about 9:40 pm. The white hatchback KA 01 AB 1234 entered the junction when its light was amber.

The blue van, driven by Suresh Nair, moved off before its light turned green and struck the hatchback."""


def structured(**overrides) -> dict:
    output = {
        "summary": "A witness saw the van move off early and strike the hatchback.",
        "persons": [
            {
                "ref": "P1",
                "name": "Mrs. Anita Desai",
                "role": "witness",
                "age": None,
                "occupation": "",
                "condition": None,
                "contact": "",
                "notes": "",
                "passage": "Statement of Anita Desai, recorded 4 October 2026.",
            },
            {
                "ref": "P2",
                "name": "Suresh Nair",
                "role": "other",
                "age": 41,
                "occupation": "Delivery driver",
                "condition": "uninjured",
                "contact": "",
                "notes": "Drove the van.",
                "passage": "The blue van, driven by Suresh Nair,",
            },
        ],
        "objects": [
            {
                "ref": "O1",
                "kind": "vehicle",
                "label": "White hatchback",
                "plate": "KA 01 AB 1234",
                "number": "",
                "make": "",
                "colour": "White",
                "passage": "The white hatchback KA 01 AB 1234 entered the junction when its light was amber.",
            },
            {
                "ref": "O2",
                "kind": "vehicle",
                "label": "Blue van",
                "plate": "",
                "number": "",
                "make": "",
                "colour": "Blue",
                "passage": "The blue van, driven by Suresh Nair,",
            },
        ],
        "locations": [
            {
                "ref": "L1",
                "kind": "incident_site",
                "address": "Ring Road and 5th Avenue junction",
                "passage": "I was waiting at the Ring Road and 5th Avenue junction at about 9:40 pm.",
            }
        ],
        "events": [
            {
                "ref": "E1",
                "category": "collision",
                "description": "The van strikes the hatchback.",
                "time_start": "2026-10-03T21:40:00+05:30",
                "time_end": None,
                "time_precision": "approximate",
                "certainty": "hypothesis",
                "location_ref": "L1",
                "note": "The time is the witness's estimate.",
                "passage": "moved off before its light turned green and struck the hatchback.",
            }
        ],
        "relationships": [
            {"from_ref": "P2", "to_ref": "O2", "kind": "drove", "label": "", "passage": "driven by Suresh Nair"},
            # Events aren't graph nodes, so this one is dropped.
            {"from_ref": "P1", "to_ref": "E1", "kind": "witnessed", "label": "Saw", "passage": "I was waiting"},
        ],
        "observations": [
            {
                "subject_ref": "O1",
                "statement": "Anita Desai says the hatchback entered the junction on amber.",
                "observed_at": "2026-10-03T21:40:00+05:30",
                "passage": "The white hatchback KA 01 AB 1234 entered the junction when its light was amber.",
            },
            {
                "subject_ref": "evidence",
                "statement": "The statement was recorded on 4 October 2026.",
                "observed_at": "2026-10-04",
                "passage": "Statement of Anita Desai, recorded 4 October 2026.",
            },
        ],
    }
    return output | overrides


class FakeModels:
    def __init__(self, client):
        self.client = client
        self.calls = []

    def generate_content(self, *, model, contents, config=None):
        self.calls.append(SimpleNamespace(model=model, contents=contents, config=config))
        if self.client.on_call:
            self.client.on_call()
        if config is not None and config.response_schema is not None:
            if self.client.structure_error:
                raise self.client.structure_error
            output = self.client.structured
            return SimpleNamespace(text=output if isinstance(output, str) else json.dumps(output))
        return SimpleNamespace(text=self.client.text)


class FakeFiles:
    def __init__(self):
        self.uploaded = []
        self.deleted = []
        self.polls = 0

    def upload(self, *, file, config=None):
        self.uploaded.append((file, config.mime_type))
        return types.File(name="files/1", state=types.FileState.PROCESSING)

    def get(self, *, name):
        self.polls += 1
        return types.File(name=name, state=types.FileState.ACTIVE)

    def delete(self, *, name):
        self.deleted.append(name)


class FakeClient:
    def __init__(self, structured=None, text="", structure_error=None, on_call=None):
        self.structured = structured if structured is not None else globals()["structured"]()
        self.text = text
        self.structure_error = structure_error
        self.on_call = on_call
        self.models = FakeModels(self)
        self.files = FakeFiles()

    def structure_calls(self):
        return [call for call in self.models.calls if call.config is not None and call.config.response_schema]

    def plain_calls(self):
        return [call for call in self.models.calls if call.config is None or not call.config.response_schema]


@pytest.fixture
def gemini(monkeypatch):
    """Turns Gemini on for the pipeline with a fake client; returns a factory."""

    def use(**kwargs) -> FakeClient:
        client = FakeClient(**kwargs)
        monkeypatch.setattr(pipeline, "get_client", lambda: client)
        return client

    return use


def stored_evidence(db, case, filename: str, data: bytes, mime_type: str, **fields) -> Evidence:
    stored = save_bytes(case.id, filename, data)
    values = {
        "title": filename,
        "file_path": stored.file_path,
        "size_bytes": stored.size_bytes,
        "sha256": stored.sha256,
        "mime_type": mime_type,
        "processing_status": "queued",
    }
    return make_evidence(db, case, **(values | fields))


def pdf_bytes(pages: list[str]) -> bytes:
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    return document.tobytes()


def docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("Medical report for Rohan Kapoor.")
    document.add_paragraph("")
    document.add_paragraph("Fractured right clavicle.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Blood alcohol"
    table.rows[0].cells[1].text = "None detected"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def reload(db, evidence: Evidence) -> Evidence:
    db.expire_all()
    return db.get(Evidence, evidence.id)


def chunks_of(db, evidence: Evidence) -> list[EvidenceChunk]:
    return list(db.scalars(select(EvidenceChunk).where(EvidenceChunk.evidence_id == evidence.id).order_by(EvidenceChunk.ordinal)))


def links_of(db, evidence: Evidence) -> list[EvidenceLink]:
    return list(db.scalars(select(EvidenceLink).where(EvidenceLink.evidence_id == evidence.id)))


def ai_rows(db, model) -> list:
    return list(db.scalars(select(model).where(model.origin == "ai").order_by(model.id)))


# --- Without Gemini --------------------------------------------------------------


def test_text_file_is_extracted_and_indexed_without_gemini(db):
    case = make_case(db)
    evidence = stored_evidence(db, case, "statement.txt", STATEMENT.encode(), "text/plain", summary="Processing failed: x")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.processed_at is not None
    assert evidence.extracted_text == STATEMENT
    assert evidence.summary is None
    chunks = chunks_of(db, evidence)
    assert chunks and all(chunk.page is None for chunk in chunks)
    assert "Suresh Nair" in " ".join(chunk.text for chunk in chunks)
    assert db.scalars(select(AiRun)).all() == []
    assert ai_rows(db, Person) == [] and ai_rows(db, Observation) == []


def test_text_encodings_and_line_endings(db):
    case = make_case(db)
    windows = stored_evidence(db, case, "notes.txt", "Café notes\r\n\r\nSecond".encode("cp1252"), "text/plain")
    utf16 = stored_evidence(db, case, "log.csv", "time,event\n21:40,impact".encode("utf-16"), "application/octet-stream")

    process_evidence(windows.id)
    process_evidence(utf16.id)

    assert reload(db, windows).extracted_text == "Café notes\n\nSecond"
    assert reload(db, utf16).extracted_text == "time,event\n21:40,impact"


def test_pdf_is_read_page_by_page(db):
    case = make_case(db)
    evidence = stored_evidence(db, case, "report.pdf", pdf_bytes(["Page one text.", "", "Page three text."]), "application/pdf")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.extracted_text == "Page one text.\n\nPage three text."
    # The blank page has no text layer, and there is no Gemini for OCR.
    assert {chunk.page for chunk in chunks_of(db, evidence)} == {1, 3}


def test_word_file_reads_paragraphs_and_tables(db):
    case = make_case(db)
    evidence = stored_evidence(
        db, case, "medical.docx", docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.extracted_text == (
        "Medical report for Rohan Kapoor.\n\nFractured right clavicle.\n\nBlood alcohol | None detected"
    )


@pytest.mark.parametrize(
    ("filename", "mime_type"),
    [("scene.png", "image/png"), ("call.wav", "audio/wav"), ("dashcam.mp4", "video/mp4"), ("dump.bin", "")],
)
def test_files_needing_gemini_finish_done_without_text(db, filename, mime_type):
    case = make_case(db)
    evidence = stored_evidence(db, case, filename, b"\x00\x01binary", mime_type)

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.extracted_text is None
    assert evidence.summary is None
    assert chunks_of(db, evidence) == []


@pytest.mark.parametrize(
    ("filename", "data", "mime_type"),
    [("call.wav", b"RIFF audio", "audio/wav"), ("scan.pdf", pdf_bytes([""]), "application/pdf")],
)
def test_reprocessing_without_gemini_keeps_text_it_cannot_read(db, filename, data, mime_type):
    case = make_case(db)
    evidence = stored_evidence(
        db, case, filename, data, mime_type, extracted_text="Caller: a van hit a car.", summary="An emergency call."
    )
    [chunk] = index_evidence(db, evidence.id, ["Caller: a van hit a car."])
    db.commit()

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert (evidence.extracted_text, evidence.summary) == ("Caller: a van hit a car.", "An emergency call.")
    assert [c.id for c in chunks_of(db, evidence)] == [chunk.id]


def test_reprocessing_without_gemini_keeps_the_summary(db):
    case = make_case(db)
    evidence = stored_evidence(db, case, "statement.txt", STATEMENT.encode(), "text/plain", summary="A witness account.")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert (evidence.processing_status, evidence.summary) == ("done", "A witness account.")
    assert evidence.extracted_text == STATEMENT


def test_evidence_without_a_file_finishes_done(db):
    case = make_case(db)
    evidence = make_evidence(db, case, processing_status="queued", extracted_text="Typed in by hand.")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.extracted_text == "Typed in by hand."


def test_missing_file_fails_with_the_reason(db):
    case = make_case(db)
    evidence = stored_evidence(db, case, "statement.txt", b"text", "text/plain")
    absolute_path(evidence.file_path).unlink()

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "failed"
    assert evidence.summary == "Processing failed: the stored file is missing"
    assert evidence.processed_at is not None


def test_corrupt_pdf_fails(db):
    case = make_case(db)
    evidence = stored_evidence(db, case, "report.pdf", b"not really a pdf", "application/pdf")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "failed"
    assert evidence.summary.startswith("Processing failed: ")
    assert evidence.processed_at is not None


def test_unknown_evidence_is_ignored():
    process_evidence(999)


# --- With Gemini -----------------------------------------------------------------


def test_status_is_processing_while_gemini_runs(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "statement.txt", STATEMENT.encode(), "text/plain")
    seen = []

    def check_status():
        db.expire_all()
        seen.append(db.get(Evidence, evidence.id).processing_status)

    gemini(on_call=check_status)
    process_evidence(evidence.id)

    assert seen == ["processing"]
    assert reload(db, evidence).processing_status == "done"


def test_structure_saves_pending_proposals_linked_to_their_passages(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "statement.txt", STATEMENT.encode(), "text/plain", type="statement")
    client = gemini()

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.summary == "A witness saw the van move off early and strike the hatchback."

    [call] = client.structure_calls()
    assert call.model == "gemini-2.5-flash"
    assert call.config.response_mime_type == "application/json"
    assert STATEMENT in call.contents[0]
    assert "Ring Road junction collision" in call.contents[0]

    [run] = db.scalars(select(AiRun)).all()
    assert (run.task, run.evidence_id, run.case_id, run.model) == ("structure", evidence.id, case.id, "gemini-2.5-flash")
    assert run.output == structured()

    persons = ai_rows(db, Person)
    assert [(p.name, p.role, p.age, p.occupation, p.condition) for p in persons] == [
        ("Mrs. Anita Desai", "witness", None, "", None),
        ("Suresh Nair", "other", 41, "Delivery driver", "uninjured"),
    ]
    objects = ai_rows(db, CaseObject)
    assert [(o.label, o.details) for o in objects] == [
        ("White hatchback", {"plate": "KA 01 AB 1234", "colour": "White"}),
        ("Blue van", {"colour": "Blue"}),
    ]
    [location] = ai_rows(db, Location)
    assert (location.kind, location.address) == ("incident_site", "Ring Road and 5th Avenue junction")
    [event] = ai_rows(db, Event)
    assert event.location_id == location.id
    assert event.time_start.isoformat() == "2026-10-03T16:10:00+00:00"
    assert (event.time_precision, event.certainty, event.ai_note) == (
        "approximate",
        "hypothesis",
        "The time is the witness's estimate.",
    )
    [relationship] = ai_rows(db, Relationship)
    assert (relationship.from_type, relationship.from_id, relationship.to_type, relationship.to_id) == (
        "person",
        persons[1].id,
        "object",
        objects[1].id,
    )
    assert (relationship.kind, relationship.label) == ("drove", "Drove")
    observations = ai_rows(db, Observation)
    assert [(o.subject_type, o.subject_id) for o in observations] == [("object", objects[0].id), ("evidence", evidence.id)]
    assert observations[1].observed_at.isoformat() == "2026-10-04T00:00:00+00:00"

    rows = [*persons, *objects, location, event, relationship, *observations]
    assert all(row.case_id == case.id and (row.origin, row.review_status) == ("ai", "pending") for row in rows)

    links = {(link.target_type, link.target_id): link for link in links_of(db, evidence)}
    expected = {("person", p.id) for p in persons} | {("object", o.id) for o in objects}
    expected |= {("location", location.id), ("event", event.id), ("relationship", relationship.id)}
    expected |= {("observation", o.id) for o in observations}
    assert set(links) == expected
    chunks = {chunk.id: chunk.text for chunk in chunks_of(db, evidence)}
    for link in links.values():
        assert link.passage
        assert link.passage in chunks[link.chunk_id]


def test_existing_people_and_vehicles_are_matched_not_duplicated(db, gemini):
    case = make_case(db)
    witness = make_person(db, case, name="Anita Desai")
    hatchback = make_object(db, case, label="Hatchback", details={"plate": "KA-01-AB-1234"})
    other_case = make_case(db, title="Another case")
    make_person(db, other_case, name="Suresh Nair")
    evidence = stored_evidence(db, case, "statement.txt", STATEMENT.encode(), "text/plain")
    client = gemini()

    process_evidence(evidence.id)

    prompt = client.structure_calls()[0].contents[0]
    assert "Anita Desai (witness)" in prompt and "plate KA-01-AB-1234" in prompt
    persons = db.scalars(select(Person).where(Person.case_id == case.id).order_by(Person.id)).all()
    assert [p.name for p in persons] == ["Anita Desai", "Suresh Nair"]
    objects = db.scalars(select(CaseObject).where(CaseObject.case_id == case.id).order_by(CaseObject.id)).all()
    assert [o.label for o in objects] == ["Hatchback", "Blue van"]
    db.refresh(witness)
    db.refresh(hatchback)
    assert (witness.origin, witness.review_status) == ("investigator", "confirmed")
    targets = {(link.target_type, link.target_id) for link in links_of(db, evidence)}
    assert {("person", witness.id), ("object", hatchback.id)} <= targets
    [observation] = db.scalars(select(Observation).where(Observation.subject_type == "object")).all()
    assert observation.subject_id == hatchback.id


def test_scanned_pdf_pages_go_to_gemini_ocr(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "report.pdf", pdf_bytes(["Typed first page.", ""]), "application/pdf")
    client = gemini(text="Handwritten second page.")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.extracted_text == "Typed first page.\n\nHandwritten second page."
    [ocr] = client.plain_calls()
    image = ocr.contents[0]
    assert image.inline_data.mime_type == "image/png"
    assert image.inline_data.data.startswith(b"\x89PNG")
    assert [(chunk.page, chunk.text) for chunk in chunks_of(db, evidence) if chunk.page == 2] == [
        (2, "Handwritten second page.")
    ]
    runs = db.scalars(select(AiRun).order_by(AiRun.id)).all()
    assert [(run.task, run.output) for run in runs[:1]] == [("ocr", {"page": 2, "text": "Handwritten second page."})]
    assert [run.task for run in runs] == ["ocr", "structure"]


def test_images_are_read_with_gemini_ocr(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "note.jpg", b"jpeg bytes", "image/jpeg")
    client = gemini(text="Licence plate KA 05 CD 5678")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.extracted_text == "Licence plate KA 05 CD 5678"
    [ocr] = client.plain_calls()
    assert ocr.contents[0].inline_data.mime_type == "image/jpeg"
    assert ocr.contents[0].inline_data.data == b"jpeg bytes"
    assert [run.task for run in db.scalars(select(AiRun).order_by(AiRun.id))] == ["ocr", "structure"]


def test_images_gemini_cannot_read_are_rendered_to_png(db, gemini):
    case = make_case(db)
    svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="200" height="80"><text x="10" y="40">KA 01</text></svg>'
    evidence = stored_evidence(db, case, "photo.svg", svg, "image/svg+xml")
    client = gemini(text="KA 01")

    process_evidence(evidence.id)

    assert reload(db, evidence).extracted_text == "KA 01"
    [ocr] = client.plain_calls()
    assert ocr.contents[0].inline_data.mime_type == "image/png"
    assert ocr.contents[0].inline_data.data.startswith(b"\x89PNG")


@pytest.mark.parametrize(
    ("filename", "mime_type", "phrase"),
    [("call.mp3", "audio/mpeg", "this recording"), ("dashcam.mp4", "video/mp4", "the soundtrack of this video")],
)
def test_audio_and_video_are_transcribed(db, gemini, filename, mime_type, phrase):
    case = make_case(db)
    evidence = stored_evidence(db, case, filename, b"media bytes", mime_type)
    client = gemini(text="[00:01] Speaker 1: There's been a crash at the junction.")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.extracted_text == "[00:01] Speaker 1: There's been a crash at the junction."
    [transcribe] = client.plain_calls()
    part, prompt = transcribe.contents
    assert part.inline_data.mime_type == mime_type
    assert phrase in prompt
    [run] = db.scalars(select(AiRun).where(AiRun.task == "transcribe")).all()
    assert run.output == {"text": evidence.extracted_text}


def test_large_media_goes_through_the_files_api(db, gemini, monkeypatch):
    monkeypatch.setattr(calls, "INLINE_LIMIT", 4)
    monkeypatch.setattr(calls, "UPLOAD_POLL_SECONDS", 0)
    case = make_case(db)
    evidence = stored_evidence(db, case, "dashcam.mp4", b"a longer video", "video/mp4")
    client = gemini(text="Speaker 1: Stop!")

    process_evidence(evidence.id)

    assert reload(db, evidence).extracted_text == "Speaker 1: Stop!"
    assert client.files.uploaded == [(absolute_path(evidence.file_path), "video/mp4")]
    assert client.files.polls == 1
    assert client.files.deleted == ["files/1"]
    [transcribe] = client.plain_calls()
    assert transcribe.contents[0].name == "files/1"


def test_media_without_text_skips_structuring(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "scene.png", b"png bytes", "image/png")
    client = gemini(text="")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "done"
    assert evidence.extracted_text is None
    assert client.structure_calls() == []


def test_gemini_error_fails_the_evidence_but_keeps_the_text(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "statement.txt", STATEMENT.encode(), "text/plain")
    gemini(structure_error=RuntimeError("429 RESOURCE_EXHAUSTED quota exceeded"))

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "failed"
    assert evidence.summary == "Processing failed: 429 RESOURCE_EXHAUSTED quota exceeded"
    assert evidence.processed_at is not None
    assert evidence.extracted_text == STATEMENT
    assert chunks_of(db, evidence)
    assert ai_rows(db, Person) == []


def test_invalid_json_fails_but_the_raw_output_is_logged(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "statement.txt", STATEMENT.encode(), "text/plain")
    gemini(structured="{not json")

    process_evidence(evidence.id)

    evidence = reload(db, evidence)
    assert evidence.processing_status == "failed"
    assert evidence.summary == "Processing failed: Gemini returned no valid JSON"
    [run] = db.scalars(select(AiRun)).all()
    assert (run.task, run.output) == ("structure", "{not json")


def test_items_that_break_the_schema_are_dropped(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "statement.txt", STATEMENT.encode(), "text/plain")
    output = structured()
    output["persons"][0]["role"] = "bystander"
    output["events"][0]["time_start"] = "about 9:40 pm"
    del output["relationships"]
    gemini(structured=output)

    process_evidence(evidence.id)

    assert reload(db, evidence).processing_status == "done"
    assert [p.name for p in ai_rows(db, Person)] == ["Suresh Nair"]
    [event] = ai_rows(db, Event)
    assert event.time_start is None
    assert ai_rows(db, Relationship) == []
    # The observation about the evidence itself still links to the evidence.
    assert len(ai_rows(db, Observation)) == 2


def test_reprocessing_replaces_pending_proposals_and_keeps_reviewed_ones(db, gemini):
    case = make_case(db)
    evidence = stored_evidence(db, case, "statement.txt", STATEMENT.encode(), "text/plain")
    gemini()
    process_evidence(evidence.id)
    suresh = db.scalar(select(Person).where(Person.name == "Suresh Nair"))
    suresh.review_status = "confirmed"
    db.commit()

    output = structured()
    output["events"][0]["description"] = "The van hits the hatchback."
    gemini(structured=output)
    process_evidence(evidence.id)

    db.expire_all()
    assert [p.name for p in db.scalars(select(Person).order_by(Person.id))] == ["Suresh Nair", "Mrs. Anita Desai"]
    assert [e.description for e in db.scalars(select(Event))] == ["The van hits the hatchback."]
    assert len(db.scalars(select(Relationship)).all()) == 1
    assert len(db.scalars(select(Observation)).all()) == 2
    links = links_of(db, evidence)
    targets = [(link.target_type, link.target_id) for link in links]
    assert len(targets) == len(set(targets))
    assert ("person", suresh.id) in targets
    chunk_ids = {chunk.id for chunk in chunks_of(db, evidence)}
    assert all(link.chunk_id in chunk_ids for link in links)


def test_proposals_other_evidence_supports_are_kept_on_reprocess(db, gemini):
    case = make_case(db)
    first = stored_evidence(db, case, "statement.txt", STATEMENT.encode(), "text/plain")
    second = stored_evidence(db, case, "statement-copy.txt", STATEMENT.encode(), "text/plain")
    gemini()
    process_evidence(first.id)
    process_evidence(second.id)
    process_evidence(first.id)

    db.expire_all()
    assert [p.name for p in db.scalars(select(Person).order_by(Person.id))] == ["Mrs. Anita Desai", "Suresh Nair"]
    assert len(db.scalars(select(Event)).all()) == 1


def test_passage_index_matches_loosely():
    chunks = [
        EvidenceChunk(id=1, text="The van, driven by Suresh   Nair, moved off."),
        EvidenceChunk(id=2, text="The hatchback entered the junction on amber and stopped."),
    ]
    index = PassageIndex(chunks)
    assert index.chunk_for("driven by suresh nair") == 1
    assert index.chunk_for("The hatchback entered the junction on amber, then stopped") == 2
    assert index.chunk_for("Nothing like it at all") is None
    assert index.chunk_for("") is None
