"""Text extraction (step 2 in docs/PLAN.md).

Local first: PyMuPDF reads PDFs page by page, python-docx reads Word files and
plain text is read as is. Gemini does what can't be read locally: OCR of images
and of PDF pages with no text layer, and transcription of audio and of the
soundtrack of video. Without a client those steps are skipped, so the result
simply has no text for them.
"""

import codecs
import mimetypes
from dataclasses import dataclass, field
from pathlib import Path

import docx
import pymupdf
from google.genai import types

from app.services.ai.calls import generate_from_file, generate_text

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_MIMES = {"application/json", "application/xml", "application/x-subrip"}
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".log", ".xml", ".srt", ".vtt"}
# Image types Gemini reads directly; others (TIFF, BMP, GIF...) are converted to PNG.
GEMINI_IMAGE_MIMES = {"image/png", "image/jpeg", "image/webp", "image/heic", "image/heif"}
OCR_DPI = 200

OCR_PROMPT = (
    "Transcribe all legible text in this image exactly as written, in reading order. "
    "Keep line breaks, and separate paragraphs with a blank line. Do not describe the image, "
    "translate, correct or add anything. If there is no legible text, return an empty response."
)
TRANSCRIBE_PROMPT = (
    "Transcribe the speech in {source} verbatim, in the language spoken. Start a new paragraph, "
    "separated by a blank line, at each change of speaker, beginning with the time it starts as "
    "[mm:ss] and a speaker label such as 'Speaker 1:', or the speaker's name if it is said. "
    "Write [inaudible] for words you cannot make out. Do not summarise, and do not describe sounds "
    "other than speech. If there is no speech, return an empty response."
)


@dataclass
class Extraction:
    """`pages` holds one string per page for PDFs (`paged`), otherwise a single
    string. `runs` are the Gemini calls made, as (task, output) for ai_runs."""

    pages: list[str] = field(default_factory=list)
    paged: bool = False
    runs: list[tuple[str, dict]] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(page.strip() for page in self.pages if page.strip())


def file_kind(path: Path, mime_type: str | None) -> tuple[str, str]:
    """Returns (kind, mime type), where kind is pdf, docx, text, image, audio, video
    or other. Falls back to the file name when the upload had no useful MIME type."""
    mime = (mime_type or "").split(";")[0].strip().lower()
    if not mime or mime == "application/octet-stream":
        mime = mimetypes.guess_type(path.name)[0] or mime
    suffix = path.suffix.lower()
    if mime == "application/pdf" or suffix == ".pdf":
        return "pdf", "application/pdf"
    if mime == DOCX_MIME or suffix == ".docx":
        return "docx", DOCX_MIME
    if mime.startswith("text/") or mime in TEXT_MIMES or suffix in TEXT_SUFFIXES:
        return "text", mime
    for kind in ("image", "audio", "video"):
        if mime.startswith(f"{kind}/"):
            return kind, mime
    return "other", mime


def extract(path: Path, mime_type: str | None, client) -> Extraction:
    kind, mime = file_kind(path, mime_type)
    if kind == "pdf":
        return read_pdf(path, client)
    if kind == "docx":
        return Extraction(pages=[read_docx(path)])
    if kind == "text":
        return Extraction(pages=[read_text(path)])
    if client is None:
        return Extraction()
    if kind == "image":
        text = ocr_image(client, path, mime)
        return Extraction(pages=[text], runs=[("ocr", {"text": text})])
    if kind in ("audio", "video"):
        source = "this recording" if kind == "audio" else "the soundtrack of this video"
        text = generate_from_file(client, path, mime, TRANSCRIBE_PROMPT.format(source=source))
        return Extraction(pages=[text], runs=[("transcribe", {"text": text})])
    return Extraction()


def read_pdf(path: Path, client) -> Extraction:
    """One string per page, with a blank line between text blocks. Pages with no
    text layer (scans) are rendered and sent to Gemini for OCR when there is a client."""
    result = Extraction(paged=True)
    with pymupdf.open(path) as document:
        for number, page in enumerate(document, start=1):
            blocks = page.get_text("blocks")
            text = "\n\n".join(block[4].strip() for block in blocks if block[6] == 0 and block[4].strip())
            if not text and client is not None:
                png = page.get_pixmap(dpi=OCR_DPI).tobytes("png")
                text = ocr_bytes(client, png, "image/png")
                result.runs.append(("ocr", {"page": number, "text": text}))
            result.pages.append(text)
    return result


def read_docx(path: Path) -> str:
    """Paragraphs, then tables row by row, separated by blank lines."""
    document = docx.Document(str(path))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))
    return "\n\n".join(part for part in parts if part.strip(" |"))


def read_text(path: Path) -> str:
    data = path.read_bytes()
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        text = data.decode("utf-16")
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")
    return text.replace("\r\n", "\n").replace("\r", "\n")


def ocr_image(client, path: Path, mime: str) -> str:
    if mime in GEMINI_IMAGE_MIMES:
        return generate_from_file(client, path, mime, OCR_PROMPT)
    pixmap = pymupdf.Pixmap(str(path))
    if pixmap.n - pixmap.alpha > 3:  # CMYK can't be written as PNG
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pixmap)
    return ocr_bytes(client, pixmap.tobytes("png"), "image/png")


def ocr_bytes(client, data: bytes, mime: str) -> str:
    return generate_text(client, [types.Part.from_bytes(data=data, mime_type=mime), OCR_PROMPT])
